import io

import openpyxl

from app.links import extract_from_file, parse_link, parse_many

FID = "1AbCdEfGhIjKlMnOpQrStUvWxYz012345"


def test_parse_common_links():
    cases = [
        f"https://drive.google.com/file/d/{FID}/view?usp=sharing",
        f"https://drive.google.com/drive/folders/{FID}?usp=drive_link",
        f"https://drive.google.com/drive/u/1/folders/{FID}",
        f"https://drive.google.com/open?id={FID}",
        f"https://drive.google.com/uc?id={FID}&export=download",
        f"https://docs.google.com/spreadsheets/d/{FID}/edit#gid=0",
        f"https://docs.google.com/document/u/0/d/{FID}/edit",
        f"drive.google.com/file/d/{FID}/view",
        FID,
    ]
    for c in cases:
        ref = parse_link(c)
        assert ref is not None and ref.id == FID, c


def test_resource_key():
    ref = parse_link(f"https://drive.google.com/drive/folders/{FID}?resourcekey=0-abcDEF")
    assert ref.resource_key == "0-abcDEF"


def test_reject_non_drive():
    assert parse_link("https://example.com/file/d/123456789012345") is None
    assert parse_link("hello") is None


def test_parse_many_dedup_and_bad():
    text = f"https://drive.google.com/file/d/{FID}/view\nfoo bar\nLink: https://drive.google.com/open?id={FID}\n"
    refs, bad = parse_many(text)
    assert len(refs) == 1 and bad == ["foo bar"]


def test_extract_xlsx_with_hyperlink():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Tên"
    ws["A2"] = f"https://drive.google.com/drive/folders/{FID}"
    ws["B3"] = "Bấm vào đây"
    ws["B3"].hyperlink = "https://drive.google.com/file/d/1ZZZZZZZZZZZZZZZZZZZZZZZZZ/view"
    ws["C4"] = "Nguyen_Van_A_2024_report_final"  # không được nhận nhầm thành ID
    buf = io.BytesIO()
    wb.save(buf)
    out = extract_from_file("links.xlsx", buf.getvalue())
    ids = {parse_link(u).id for u in out}
    assert ids == {FID, "1ZZZZZZZZZZZZZZZZZZZZZZZZZ"}


def test_extract_csv():
    data = f"stt,link\n1,https://drive.google.com/file/d/{FID}/view\n".encode()
    assert len(extract_from_file("a.csv", data)) == 1
